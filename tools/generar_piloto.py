"""Genera Turismo-piloto.xlsx: 100 establecimientos ficticios con datos de prueba.

Estructura fija de cada pestaña de establecimiento:
Año | Mes | País | Provincia | Ciudad | Cantidad | Motivo de visita | Edad | Género
Pestañas de sistema (el dashboard las ignora): _Instrucciones, _Catalogos, Plantilla.
Semilla fija: mismos datos en cada ejecución.
"""
import random
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

random.seed(20260929)
FILAS_VALIDADAS = 3000
COLUMNAS = ["Año", "Mes", "País", "Provincia", "Ciudad", "Cantidad", "Motivo de visita", "Edad", "Género"]
ANIOS = [2025, 2026]
MESES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto",
         "Septiembre", "Octubre", "Noviembre", "Diciembre"]
ULTIMO_MES_2026 = 9
MOTIVOS = ["Turismo", "Gastronomía", "Descanso", "Negocios", "Visita familiar", "Evento"]
GENEROS = ["Masculino", "Femenino"]
# Ciudad, Provincia, Lat, Lon, peso de procedencia
CIUDADES = [
    ("Ambato", "Tungurahua", -1.2491, -78.6168, 30), ("Quito", "Pichincha", -0.1807, -78.4678, 22),
    ("Riobamba", "Chimborazo", -1.6636, -78.6546, 10), ("Latacunga", "Cotopaxi", -0.9352, -78.6155, 9),
    ("Guayaquil", "Guayas", -2.1894, -79.8891, 8), ("Cuenca", "Azuay", -2.9001, -79.0059, 5),
    ("Baños", "Tungurahua", -1.3964, -78.4247, 5), ("Pelileo", "Tungurahua", -1.3306, -78.5436, 5),
    ("Píllaro", "Tungurahua", -1.1667, -78.5333, 4), ("Patate", "Tungurahua", -1.3150, -78.5100, 4),
    ("Puyo", "Pastaza", -1.4924, -78.0024, 3), ("Ibarra", "Imbabura", -0.3517, -78.1223, 3),
    ("Santo Domingo", "Santo Domingo de los Tsáchilas", -0.2530, -79.1754, 3),
    ("Manta", "Manabí", -0.9677, -80.7089, 2), ("Portoviejo", "Manabí", -1.0546, -80.4545, 2),
    ("Loja", "Loja", -3.9931, -79.2042, 2), ("Machala", "El Oro", -3.2581, -79.9554, 2),
    ("Tena", "Napo", -0.9938, -77.8129, 2), ("Guaranda", "Bolívar", -1.5926, -79.0010, 2),
    ("Esmeraldas", "Esmeraldas", 0.9682, -79.6517, 1), ("Tulcán", "Carchi", 0.8119, -77.7173, 1),
    ("Babahoyo", "Los Ríos", -1.8022, -79.5344, 1), ("Macas", "Morona Santiago", -2.3087, -78.1114, 1),
    ("Azogues", "Cañar", -2.7397, -78.8486, 1),
]
# País, Lat, Lon, peso entre extranjeros
PAISES = [
    ("Ecuador", -1.8312, -78.1834, 0), ("Colombia", 4.5709, -74.2973, 25), ("Estados Unidos", 37.0902, -95.7129, 20),
    ("Perú", -9.1900, -75.0152, 12), ("España", 40.4637, -3.7492, 10), ("Alemania", 51.1657, 10.4515, 8),
    ("Francia", 46.2276, 2.2137, 6), ("Venezuela", 6.4238, -66.5897, 6), ("Argentina", -38.4161, -63.6167, 5),
    ("Chile", -35.6751, -71.5430, 4), ("Canadá", 56.1304, -106.3468, 4),
]
TIPOS = [("Restaurante", 30), ("Heladería", 12), ("Hostería", 18), ("Hacienda", 10), ("Cafetería", 12),
         ("Parque", 6), ("Mirador", 6), ("Hostal", 6)]
NOMBRES = ["Los Andes", "El Mirador", "Valle Dorado", "La Moya", "Río Patate", "Las Nubes", "El Tungurahua",
           "San Cristóbal", "Los Nogales", "La Colina", "El Huerto", "Las Mandarinas", "Sol del Valle",
           "La Vertiente", "El Chaquiñán", "Los Capulíes", "Mama Juana", "El Trapiche", "Las Orquídeas",
           "El Colibrí", "Pucará", "Leito", "Sucre", "Los Guabos", "El Aguacate", "La Cascada",
           "Las Palmas", "El Rosal", "Quinta Aurora", "El Nogal", "Tierra Viva", "Cumbres", "El Batán",
           "La Hacienda Vieja", "Los Arrayanes"]
PESO_MES = [0.8, 1.3, 1.0, 1.1, 0.8, 0.9, 1.3, 1.4, 0.9, 0.9, 1.1, 1.5]  # Carnaval, vacaciones sierra, fin de año


def nombres_establecimientos(n):
    usados, lista = set(), []
    while len(lista) < n:
        tipo = random.choices([t for t, _ in TIPOS], [p for _, p in TIPOS])[0]
        nombre = f"{tipo} {random.choice(NOMBRES)}"[:31]
        if nombre not in usados:
            usados.add(nombre)
            lista.append(nombre)
    return lista


def fila_visitantes(anio, mes, tipo):
    extranjero = random.random() < 0.18
    if extranjero:
        p = random.choices(PAISES[1:], [x[3] for x in PAISES[1:]])[0]
        pais, prov, ciudad = p[0], "", ""
    else:
        c = random.choices(CIUDADES, [x[4] for x in CIUDADES])[0]
        pais, prov, ciudad = "Ecuador", c[1], c[0]
    motivos_peso = {"Restaurante": [3, 6, 1, 1, 2, 1], "Heladería": [4, 5, 1, 0.3, 2, 0.5],
                    "Hostería": [5, 2, 5, 1, 1, 1], "Hacienda": [4, 1, 4, 0.5, 1, 3]}.get(tipo, [6, 1, 2, 0.5, 2, 1])
    motivo = random.choices(MOTIVOS, motivos_peso)[0]
    edad = max(5, min(85, int(random.gauss(36, 13))))
    cantidad = random.choices([1, 2, 3, 4, 5, 6, 8, 10, 15], [20, 30, 15, 15, 8, 5, 3, 3, 1])[0]
    return [anio, MESES[mes - 1], pais, prov, ciudad, cantidad, motivo, edad, random.choice(GENEROS)]


ENC_FILL = PatternFill("solid", start_color="30A848", end_color="30A848")
ENC_FONT = Font(bold=True, color="FFFFFF")


def preparar_hoja(ws, n_cat):
    ws.append(COLUMNAS)
    for c in ws[1]:
        c.fill, c.font, c.alignment = ENC_FILL, ENC_FONT, Alignment(horizontal="center")
    for col, ancho in zip("ABCDEFGHI", [8, 12, 16, 26, 16, 10, 18, 8, 12]):
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = "A2"
    fin = FILAS_VALIDADAS

    def lista(col, rango, titulo):
        dv = DataValidation(type="list", formula1=rango, allow_blank=True, showErrorMessage=True,
                            errorTitle="Valor no válido", error=f"Elija un valor de la lista de {titulo}.")
        dv.add(f"{col}2:{col}{fin}")
        ws.add_data_validation(dv)

    lista("A", "_Catalogos!$A$2:$A$5", "años")
    lista("B", "_Catalogos!$B$2:$B$13", "meses")
    lista("C", f"_Catalogos!$C$2:$C${1 + n_cat['pais']}", "países")
    lista("D", f"_Catalogos!$G$2:$G${1 + n_cat['prov']}", "provincias")
    lista("E", f"_Catalogos!$D$2:$D${1 + n_cat['ciudad']}", "ciudades")
    lista("G", f"_Catalogos!$K$2:$K${1 + len(MOTIVOS)}", "motivos")
    lista("I", f"_Catalogos!$L$2:$L${1 + len(GENEROS)}", "géneros")
    for col, lo, hi, msg in [("F", 1, 500, "La cantidad es un entero entre 1 y 500."),
                             ("H", 0, 110, "La edad es un entero entre 0 y 110.")]:
        dv = DataValidation(type="whole", operator="between", formula1=str(lo), formula2=str(hi),
                            allow_blank=True, showErrorMessage=True, errorTitle="Número no válido", error=msg)
        dv.add(f"{col}2:{col}{fin}")
        ws.add_data_validation(dv)


def main():
    wb = Workbook()
    ins = wb.active
    ins.title = "_Instrucciones"
    for linea in [
        "INSTRUCCIONES — Estadísticas de visitantes, Cantón Patate (DATOS DE PRUEBA)",
        "",
        "1. Cada pestaña es un establecimiento. El nombre de la pestaña es el nombre que aparece en el dashboard.",
        "2. Para agregar un establecimiento: clic derecho en 'Plantilla' > Duplicar, y renombre la copia.",
        "3. Cada fila es un grupo de visitantes del mismo mes, procedencia, motivo, edad y género.",
        "   Una persona sola: Cantidad = 1.",
        "4. Use siempre los desplegables. Visitantes extranjeros: País distinto de Ecuador y deje vacías Provincia y Ciudad.",
        "5. No mueva ni renombre las columnas. No edite las pestañas que empiezan con '_'.",
        "6. Nacional o extranjero no se escribe: el dashboard lo calcula a partir del País.",
    ]:
        ins.append([linea])
    ins["A1"].font = Font(bold=True, size=14, color="30A848")
    ins.column_dimensions["A"].width = 110

    cat = wb.create_sheet("_Catalogos")
    provincias = sorted({c[1] for c in CIUDADES} | {"Carchi", "Sucumbíos", "Orellana", "Zamora Chinchipe",
                                                     "Santa Elena", "Galápagos"})
    cols = {"A": ["Año"] + [2024, 2025, 2026, 2027], "B": ["Mes"] + MESES,
            "C": ["País"] + [p[0] for p in PAISES], "D": ["Ciudad"] + [c[0] for c in CIUDADES],
            "E": ["Ciudad_Lat"] + [c[2] for c in CIUDADES], "F": ["Ciudad_Lon"] + [c[3] for c in CIUDADES],
            "G": ["Provincia"] + provincias, "H": ["Ciudad_Provincia"] + [c[1] for c in CIUDADES],
            "I": ["Pais_Lat"] + [p[1] for p in PAISES], "J": ["Pais_Lon"] + [p[2] for p in PAISES],
            "K": ["Motivo"] + MOTIVOS, "L": ["Género"] + GENEROS}
    for col, valores in cols.items():
        for i, v in enumerate(valores, start=1):
            cat[f"{col}{i}"] = v
        cat[f"{col}1"].font = Font(bold=True)
        cat.column_dimensions[col].width = 16
    n_cat = {"pais": len(PAISES), "prov": len(provincias), "ciudad": len(CIUDADES)}

    plantilla = wb.create_sheet("Plantilla")
    preparar_hoja(plantilla, n_cat)

    total_filas = 0
    for nombre in nombres_establecimientos(100):
        ws = wb.create_sheet(nombre)
        preparar_hoja(ws, n_cat)
        tipo = nombre.split(" ")[0]
        tamano = random.choice([0.4, 0.7, 1.0, 1.0, 1.4, 2.0])
        for anio in ANIOS:
            for mes in range(1, 13):
                if anio == 2026 and mes > ULTIMO_MES_2026:
                    break
                n = max(1, int(random.gauss(12, 3) * tamano * PESO_MES[mes - 1] * (1.12 if anio == 2026 else 1)))
                for _ in range(n):
                    ws.append(fila_visitantes(anio, mes, tipo))
                    total_filas += 1
    wb.save("Turismo-piloto.xlsx")
    print(f"pestañas={len(wb.sheetnames)} establecimientos=100 filas={total_filas}")


if __name__ == "__main__":
    main()
