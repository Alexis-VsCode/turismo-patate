"""Oráculo de conciliación: calcula con openpyxl, sin compartir código con el JS,
los valores que el dashboard debe mostrar en cada escenario de test/fixtures/escenarios.json.

Uso: python tools/oraculo.py test/fixtures/piloto-publicado.xlsx [salida.json]
Escribe test/fixtures/esperado.json, o la salida indicada. Lee los dos formatos de pestaña: el anterior (Cantidad, Edad
y Género) y el mensual (mujeres y hombres por rango de edad, más personas con discapacidad).

Autor: Kevin Alexis Barrera Llerena 2026
"""
import json
import sys
import unicodedata
from collections import defaultdict
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
         "septiembre", "octubre", "noviembre", "diciembre"]
RANGOS = ["0-30", "31-45", "46-60", "61+"]


def sin_tildes(texto):
    return "".join(c for c in unicodedata.normalize("NFD", str(texto)) if unicodedata.category(c) != "Mn").strip().lower()


def rango(edad):
    for limite, nombre in [(30, "0-30"), (45, "31-45"), (60, "46-60")]:
        if edad <= limite:
            return nombre
    return "61+"


RANGOS_DISCAPACIDAD = ["1-5", "6-10", "11-15", "16+"]
COLUMNAS_MENSUALES = [(f"{sexo} {r}", genero, i) for sexo, genero in (("mujeres", "Femenino"), ("hombres", "Masculino"))
                      for i, r in enumerate(RANGOS)]


def rango_discapacidad(personas):
    return "1-5" if personas <= 5 else "6-10" if personas <= 10 else "11-15" if personas <= 15 else "16+"


def numero(valor):
    return 0 if valor is None or str(valor).strip() == "" else int(valor)


def leer_mensual(nombre, encabezado, iterador, filas, discapacidad):
    """Pestaña del formato mensual: cada número de las ocho columnas es un grupo de visitantes de ese género y rango."""
    idx = {c: encabezado.index(c) for c in ["ano", "mes", "pais", "ciudad", "motivo de visita", "personas con discapacidad"]}
    idx_sexo = {col: encabezado.index(col) for col, _, _ in COLUMNAS_MENSUALES}
    for r in iterador:
        conteos = [(genero, i, numero(r[idx_sexo[col]])) for col, genero, i in COLUMNAS_MENSUALES]
        total = sum(n for _, _, n in conteos)
        personas = numero(r[idx["personas con discapacidad"]])
        if total == 0 and personas == 0:
            continue
        nacional = sin_tildes(r[idx["pais"]]) == "ecuador"
        ciudad = str(r[idx["ciudad"]] or "").strip() if nacional else ""
        base = {
            "establecimiento": " ".join(nombre.split()), "anio": int(r[idx["ano"]]), "mes": MESES.index(sin_tildes(r[idx["mes"]])),
            "pais": str(r[idx["pais"]]).strip(), "nacional": nacional, "ciudad": ciudad,
            "provincia": PROVINCIA_DE.get(sin_tildes(ciudad), "") if nacional else "", "motivo": str(r[idx["motivo de visita"]]).strip(),
        }
        for genero, i, n in conteos:
            if n:
                filas.append({**base, "cantidad": n, "rango": RANGOS[i], "genero": genero})
        if personas:
            discapacidad.append({**base, "personas": personas})


PROVINCIA_DE = {}


def cargar_provincias(libro):
    """Provincia de cada ciudad según la pestaña _Catalogos del propio libro (columnas Ciudad y Ciudad_Provincia)."""
    if "_Catalogos" not in libro.sheetnames:
        return
    filas = libro["_Catalogos"].iter_rows(values_only=True)
    enc = [sin_tildes(v or "") for v in next(filas)]
    ci, cp = enc.index("ciudad"), enc.index("ciudad_provincia")
    for r in filas:
        r = tuple(r) + (None,) * (max(ci, cp) + 1 - len(r))      # las filas cortas de la hoja vienen recortadas
        if r[ci]:
            PROVINCIA_DE[sin_tildes(r[ci])] = str(r[cp] or "").strip()


def leer(ruta):
    libro = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    filas = []
    discapacidad = []
    cargar_provincias(libro)
    for nombre in libro.sheetnames:
        if nombre.startswith("_") or "plantilla" in sin_tildes(nombre):
            continue
        hoja = libro[nombre]
        iterador = hoja.iter_rows(values_only=True)
        encabezado = [sin_tildes(v or "") for v in next(iterador)]
        if "mujeres 0-30" in encabezado:
            leer_mensual(nombre, encabezado, iterador, filas, discapacidad)
            continue
        idx = {c: encabezado.index(c) for c in ["ano", "mes", "pais", "provincia", "ciudad", "cantidad", "motivo de visita", "edad", "genero"]}
        for r in iterador:
            if r is None or all(v is None or str(v).strip() == "" for v in r):
                continue
            nacional = sin_tildes(r[idx["pais"]]) == "ecuador"
            filas.append({
                "establecimiento": " ".join(nombre.split()),
                "anio": int(r[idx["ano"]]), "mes": MESES.index(sin_tildes(r[idx["mes"]])),
                "pais": str(r[idx["pais"]]).strip(), "nacional": nacional,
                "ciudad": str(r[idx["ciudad"]] or "").strip() if nacional else "",
                "provincia": str(r[idx["provincia"]] or "").strip() if nacional else "",
                "cantidad": int(r[idx["cantidad"]]), "motivo": str(r[idx["motivo de visita"]]).strip(),
                "rango": rango(int(r[idx["edad"]])), "genero": str(r[idx["genero"]]).strip(),
            })
    return filas, discapacidad


def cumple(f, filtros, ignorar=()):
    def activo(k):
        return k not in ignorar and filtros.get(k) not in (None, "")
    p = filtros.get("procedencia") or ""
    if activo("procedencia"):
        if p == "NAC" and not f["nacional"]: return False
        if p == "EXT" and f["nacional"]: return False
        if p.startswith("C:") and not (f["nacional"] and f["ciudad"] == p[2:]): return False
        if p.startswith("PR:") and not (f["nacional"] and f["provincia"] == p[3:]): return False
        if p.startswith("P:") and f["pais"] != p[2:]: return False
    return ((not activo("establecimiento") or f["establecimiento"] == filtros["establecimiento"]) and
            (not activo("anio") or f["anio"] == filtros["anio"]) and
            (not activo("mes") or f["mes"] == filtros["mes"]) and
            (not activo("motivo") or f["motivo"] == filtros["motivo"]) and
            (not activo("rangoEdad") or f["rango"] == filtros["rangoEdad"]) and
            (not activo("genero") or f["genero"] == filtros["genero"]))


def ordenado(sumas):
    return [{"nombre": k, "valor": v} for k, v in sorted(sumas.items(), key=lambda kv: (-kv[1], kv[0]))]


def suma_por(filas, clave):
    s = defaultdict(int)
    for f in filas:
        if clave(f):
            s[clave(f)] += f["cantidad"]
    return s


def escenario(filas, discapacidad, filtros):
    sel = [f for f in filas if cumple(f, filtros)]
    total = sum(f["cantidad"] for f in sel)
    nac = sum(f["cantidad"] for f in sel if f["nacional"])
    base_evo = [f for f in filas if cumple(f, filtros, ignorar=("anio", "mes"))]
    anio = filtros.get("anio") or max((f["anio"] for f in base_evo), default=None)
    actual, anterior = [0] * 12, [0] * 12
    for f in base_evo:
        if f["anio"] == anio: actual[f["mes"]] += f["cantidad"]
        if f["anio"] == (anio or 0) - 1: anterior[f["mes"]] += f["cantidad"]
    ultimo = max((f["mes"] for f in base_evo if f["anio"] == anio), default=-1)
    generos = ["Masculino", "Femenino"]
    eg = suma_por(sel, lambda f: f"{f['rango']}|{f['genero']}")
    mes_f = filtros.get("mes")
    ult = ultimo
    act_v = sum(f["cantidad"] for f in base_evo if f["anio"] == anio and (f["mes"] == mes_f if mes_f is not None else f["mes"] <= ult))
    ant_v = sum(f["cantidad"] for f in base_evo if f["anio"] == (anio or 0) - 1 and (f["mes"] == mes_f if mes_f is not None else f["mes"] <= ult))
    resultado = {
        "variacion": {"anio": anio, "actual": act_v, "anterior": ant_v},
        "kpis": {"total": total, "nacionales": nac, "extranjeros": total - nac},
        "evolucion": {"anio": anio, "actual": [v if i <= ultimo else None for i, v in enumerate(actual)], "anterior": anterior},
        "porMotivo": ordenado(suma_por(sel, lambda f: f["motivo"])),
        "porCiudad": ordenado(suma_por([f for f in sel if f["nacional"]], lambda f: f["ciudad"])),
        "porProvincia": ordenado(suma_por([f for f in sel if f["nacional"]], lambda f: f["provincia"])),
        "porPais": ordenado(suma_por([f for f in sel if not f["nacional"]], lambda f: f["pais"])),
        "edadGenero": [{"genero": g, "valores": [eg.get(f"{r}|{g}", 0) for r in RANGOS]} for g in generos],
    }
    if discapacidad:
        # La discapacidad no se cruza con la edad ni con el género: sus filas y los visitantes del porcentaje ignoran esos dos filtros
        ignorar = ("rangoEdad", "genero")
        regs = [d for d in discapacidad if cumple(d, filtros, ignorar)]
        visitantes = sum(f["cantidad"] for f in filas if cumple(f, filtros, ignorar))
        personas = sum(d["personas"] for d in regs)
        por_rango = {r: sum(1 for d in regs if rango_discapacidad(d["personas"]) == r) for r in RANGOS_DISCAPACIDAD}
        resultado["discapacidad"] = {
            "personas": personas, "registros": len(regs), "pctVisitantes": personas / visitantes if visitantes else None,
            "rangos": {r: n for r, n in por_rango.items()},
        }
    return resultado


def main():
    ruta = sys.argv[1] if len(sys.argv) > 1 else str(RAIZ / "test/fixtures/piloto-publicado.xlsx")
    filas, discapacidad = leer(ruta)
    escenarios = json.loads((RAIZ / "test/fixtures/escenarios.json").read_text(encoding="utf-8"))
    salida = {e["nombre"]: escenario(filas, discapacidad, e["filtros"]) for e in escenarios}
    destino = Path(sys.argv[2]) if len(sys.argv) > 2 else RAIZ / "test/fixtures/esperado.json"
    destino.write_text(json.dumps(salida, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"filas={len(filas)} escenarios={len(salida)} " + " ".join(f"{k}:{v['kpis']['total']}" for k, v in salida.items()))


if __name__ == "__main__":
    main()
