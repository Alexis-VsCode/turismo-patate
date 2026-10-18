"""Prueba de la plantilla del libro de captura: protección, celdas de entrada, columnas calculadas y validaciones.

Genera un libro pequeño (modo fixture) en un directorio temporal y lo abre con openpyxl.

Uso:
    python test/excel/plantilla.test.py

Autor: Kevin Alexis Barrera Llerena 2026
"""
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent.parent
GENERADOR = RAIZ / "tools" / "excel" / "generar-excel.py"
ENTRADA = ["A", "B", "C", "E", "F"] + list("GHIJKLMN") + ["O"]
CALCULADAS = ["D", "P", "Q", "R", "S", "T", "U"]


def rangos_validados(hoja, tipo):
    """Celdas (como texto, p. ej. 'G2') cubiertas por validaciones del tipo dado, en una muestra de filas."""
    cubiertas = set()
    for v in hoja.data_validations.dataValidation:
        if v.type == tipo:
            for rango in v.sqref.ranges:
                for fila in (rango.min_row, rango.max_row):
                    for col in range(rango.min_col, rango.max_col + 1):
                        cubiertas.add(f"{openpyxl.utils.get_column_letter(col)}{fila}")
    return cubiertas


class PlantillaEstandarizada(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        salida = Path(cls.tmp.name) / "libro.xlsx"
        r = subprocess.run([sys.executable, str(GENERADOR), "--modo", "fixture", "--salida", str(salida)],
                           capture_output=True, text=True, encoding="utf-8")
        assert r.returncode == 0, r.stderr
        cls.libro = openpyxl.load_workbook(salida)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def hojas_captura(self):
        nombres = [n for n in self.libro.sheetnames if not n.startswith("_")]
        self.assertIn("Plantilla", nombres)
        self.assertGreater(len(nombres), 1)
        return [self.libro[n] for n in nombres]

    def test_hojas_de_captura_protegidas_sin_contrasena(self):
        for hoja in self.hojas_captura():
            self.assertTrue(hoja.protection.sheet, hoja.title)
            self.assertFalse(hoja.protection.password, f"{hoja.title}: sin contrasena")

    def test_catalogos_no_se_protege(self):
        self.assertFalse(self.libro["_Catalogos"].protection.sheet)

    def test_celdas_de_entrada_desbloqueadas(self):
        for hoja in self.hojas_captura():
            for col in ENTRADA:
                for fila in (2, 3, 300, 600):
                    self.assertFalse(hoja[f"{col}{fila}"].protection.locked, f"{hoja.title}!{col}{fila}")

    def test_columnas_grises_bloqueadas_y_con_formula(self):
        for hoja in self.hojas_captura():
            for col in CALCULADAS:
                for fila in (2, 3, 300, 600):
                    celda = hoja[f"{col}{fila}"]
                    self.assertTrue(celda.protection.locked, f"{hoja.title}!{col}{fila}")
                    self.assertIsInstance(celda.value, str, f"{hoja.title}!{col}{fila}")
                    self.assertTrue(celda.value.startswith("="), f"{hoja.title}!{col}{fila}")

    def test_encabezados_bloqueados(self):
        for hoja in self.hojas_captura():
            for col in ENTRADA + CALCULADAS:
                self.assertTrue(hoja[f"{col}1"].protection.locked, f"{hoja.title}!{col}1")

    def test_validaciones_de_entrada_presentes(self):
        hoja = self.libro["Plantilla"]
        listas = rangos_validados(hoja, "list")
        for col in "ABCEF":
            self.assertIn(f"{col}2", listas)
            self.assertIn(f"{col}600", listas)
        enteros = rangos_validados(hoja, "whole")
        for col in "GHIJKLMN":
            self.assertIn(f"{col}2", enteros)
            self.assertIn(f"{col}600", enteros)
        self.assertTrue({"O2", "O600"} <= rangos_validados(hoja, "custom"))

    def test_regla_discapacidad_y_enteros_no_negativos(self):
        hoja = self.libro["Plantilla"]
        for v in hoja.data_validations.dataValidation:
            sq = str(v.sqref)
            if sq.startswith("G2"):
                self.assertEqual(str(v.formula1), "0")
            if sq.startswith("O2"):
                self.assertIn("SUM(G2:N2)", v.formula1)
                self.assertIn("O2>=0", v.formula1)

    def test_formato_condicional_rojo_cuando_estado_no_es_ok(self):
        hoja = self.libro["Plantilla"]
        formulas = [r.formula[0] for rangos in hoja.conditional_formatting for r in rangos.rules if r.formula]
        self.assertTrue(any("$U2" in f and "OK" in f for f in formulas), formulas)

    def test_la_guia_no_promete_que_una_fila_en_rojo_no_se_publica(self):
        textos = " ".join(str(c.value) for fila in self.libro["_Inicio"].iter_rows() for c in fila if c.value)
        self.assertNotIn("no se publica en el tablero", textos)
        self.assertIn("panel de problemas", textos)

    def test_catalogo_ofrece_otros_paises_sin_coordenadas(self):
        hoja = self.libro["_Catalogos"]
        filas = [[c.value for c in fila] for fila in hoja.iter_rows()]
        col = filas[0].index("País")
        otros = [f for f in filas[1:] if f[col] == "Extranjeros (sin especificar)"]
        self.assertEqual(len(otros), 1)
        self.assertIsNone(otros[0][col + 1])
        self.assertIsNone(otros[0][col + 2])

    def test_la_guia_explica_que_los_extranjeros_van_en_fila_aparte(self):
        textos = " ".join(str(c.value) for fila in self.libro["_Inicio"].iter_rows() for c in fila if c.value)
        self.assertIn("fila aparte", textos)
        self.assertIn("Extranjeros (sin especificar)", textos)


class SinDatosReales(unittest.TestCase):
    def test_el_generador_no_guarda_establecimientos_ni_cifras_reales(self):
        """Los datos reales viven solo en la hoja de Google; el generador arma la plantilla y datos ficticios."""
        fuente = GENERADOR.read_text(encoding="utf-8")
        for real in ("El Valle", "Tu Heladería", "La Casta", "Patate Gardens", "Quinlata"):
            self.assertNotIn(real, fuente)


if __name__ == "__main__":
    unittest.main()
